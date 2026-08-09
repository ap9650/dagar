"""
The feedback report as an editable Word document.

    npm run feedback:doc

    → docs/deck/Dagar-Feedback-Report.docx

The PDF beside it is the one to submit: fixed layout, nothing shifts. This
exists for the case where a form wants a .docx, or where someone needs to add a
covering paragraph before sending it on.

Same source and the same rule as the PDF: it reads `feedback-export/feedback.csv`
written by `scripts/export-feedback.ts`, which owns the masking. Several
respondents are children, so there is one careful path to those rows rather than
two. Nothing is filtered out, including the criticism.
"""
import csv
import pathlib
from collections import Counter

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

SRC = pathlib.Path("feedback-export/feedback.csv")
OUT = pathlib.Path("docs/deck/Dagar-Feedback-Report.docx")

INK = RGBColor(0x1A, 0x27, 0x33)
BODY = RGBColor(0x3F, 0x4A, 0x57)
MUTED = RGBColor(0x5B, 0x66, 0x73)
PRIMARY = RGBColor(0x0F, 0x76, 0x6E)
AMBER = RGBColor(0xB4, 0x53, 0x09)

LABELS = {
    "yes": "Yes", "a_bit": "A bit", "no": "No", "not_really": "Not really",
    "student": "Student", "parent": "Parent", "teacher": "Teacher",
    "other": "Other", "chapters": "More chapters",
    "practice": "More practice questions", "tutor": "A better tutor",
}


def label(value):
    return LABELS.get((value or "").strip(), (value or "—").strip() or "—")


def para(doc, text, size=10, colour=BODY, bold=False, after=6, before=0,
         italic=False, indent=0.0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = colour
    run.bold = bold
    run.italic = italic
    run.font.name = "Calibri"
    return p


def bar_row(table, name, n, total, colour):
    """One result as a row: label, a bar drawn in block characters, the count."""
    cells = table.add_row().cells
    for cell, text, size, col, bold, align in (
        (cells[0], name, 10, BODY, False, WD_ALIGN_PARAGRAPH.LEFT),
        (cells[1], "█" * round(24 * n / max(total, 1)), 10, colour, False,
         WD_ALIGN_PARAGRAPH.LEFT),
        (cells[2], str(n), 10, colour, True, WD_ALIGN_PARAGRAPH.RIGHT),
    ):
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.font.color.rgb = col
        run.bold = bold
        run.font.name = "Calibri"


def block(doc, title, counts, order, total, colours=None):
    para(doc, title, size=11, colour=INK, bold=True, after=4, before=8)
    table = doc.add_table(rows=0, cols=3)
    table.autofit = False
    for row in (order if order else [k for k, _ in counts.most_common()]):
        colour = (colours or {}).get(row, PRIMARY)
        bar_row(table, label(row), counts.get(row, 0), total, colour)
    for r in table.rows:
        r.cells[0].width = Inches(2.4)
        r.cells[1].width = Inches(2.9)
        r.cells[2].width = Inches(0.5)


def main():
    if not SRC.exists():
        raise SystemExit("run `npm run feedback:export` first")
    rows = list(csv.DictReader(SRC.open()))
    total = len(rows)

    doc = Document()
    # Setting the font per run is not enough on its own: Word falls back to the
    # Normal style, and the preview rendered everything in Times. The style has
    # to carry the name too, including the east-asian slot.
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.8)
        section.left_margin = section.right_margin = Inches(0.9)

    para(doc, "DAGAR", size=8, colour=PRIMARY, bold=True, after=2)
    para(doc, "What our users told us", size=22, colour=INK, bold=True, after=8)

    para(doc,
         f"Every one of the {total} responses below was left inside the product "
         "by a signed-in account that had used it. Nobody is reacting to a "
         "screenshot or a demo. Spelling and grammar are exactly as they were "
         "typed, and nothing has been left out, including the criticism.",
         size=10.5, colour=BODY, after=6)
    para(doc,
         "No names or email addresses appear in this report. Several "
         "respondents are children aged 11 to 14 and under the DPDP Act 2023 "
         "they are minors, so each row carries only a truncated reference, "
         "enough to show the responses came from distinct people.",
         size=9.5, colour=MUTED, after=12)

    understood = Counter(x["understood"] for x in rows)
    again = Counter(x["would_use_again"] for x in rows)
    improve = Counter(x["improve_most"] for x in rows)
    who = Counter(x["respondent"] for x in rows)
    helped = understood["yes"] + understood["a_bit"]

    para(doc, "The numbers", size=14, colour=INK, bold=True, after=4, before=6)
    para(doc,
         f"{helped} of {total} said Dagar helped them understand something. "
         f"{again['yes']} of {total} would use it again.",
         size=10.5, colour=INK, after=4)

    block(doc, "Did it help you understand something?", understood,
          ["yes", "a_bit", "no"], total,
          {"a_bit": AMBER, "no": AMBER})
    block(doc, "Would you use Dagar again?", again, ["yes", "no"], total,
          {"no": AMBER})
    block(doc, "If we could do only ONE more thing", improve, None, total)
    block(doc, "Who answered", who, None, total)

    # No page break here on purpose. Every way of inserting one left a visible
    # glyph in preview, and a stray box in the middle of a document being
    # submitted is a worse outcome than the section simply flowing on. The
    # heading is enough of a divider, and whoever edits this can add a break
    # themselves in one keystroke.
    para(doc, "Every response, in full", size=14, colour=INK, bold=True,
         after=2, before=18)
    para(doc, "Ordered as collected. A dash means the question was left blank.",
         size=9.5, colour=MUTED, after=10)

    for i, row in enumerate(rows, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        head = p.add_run(f"{i:02d}   {label(row['respondent']).upper()}")
        head.bold = True
        head.font.size = Pt(10)
        head.font.color.rgb = PRIMARY
        head.font.name = "Calibri"
        meta = p.add_run(
            f"      {row['date']}  ·  {row['language']}  ·  "
            f"ref {row['user_ref']}")
        meta.font.size = Pt(8.5)
        meta.font.color.rgb = MUTED
        meta.font.name = "Calibri"

        para(doc,
             f"Understood: {label(row['understood'])}   ·   "
             f"Would use again: {label(row['would_use_again'])}   ·   "
             f"Wants most: {label(row['improve_most'])}",
             size=9, colour=BODY, after=3)

        worked = (row.get("what_worked") or "").strip()
        confused = (row.get("what_confused") or "").strip()
        if worked:
            para(doc, f"What worked:  “{worked}”", size=10, colour=INK,
                 after=2, indent=0.2)
        if confused:
            para(doc, f"What confused them:  “{confused}”", size=10,
                 colour=INK, after=2, indent=0.2)
        if not worked and not confused:
            para(doc, "No written comment.", size=9.5, colour=MUTED, after=2,
                 indent=0.2, italic=True)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT}  ({total} responses, "
          f"{OUT.stat().st_size / 1000:.0f} KB)")


if __name__ == "__main__":
    main()
